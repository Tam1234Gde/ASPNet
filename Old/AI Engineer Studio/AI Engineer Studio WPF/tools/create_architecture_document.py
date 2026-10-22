from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent / 'AI_Engineer_Studio_Architecture_and_Agent_Guide.docx'

NAVY = '101828'
BLUE = '2563EB'
CYAN = '06B6D4'
LIGHT = 'EFF6FF'
MUTED = '475467'
WHITE = 'FFFFFF'
GREEN = '027A48'
ORANGE = 'B54708'


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    properties.append(shading)


def border_bottom(paragraph, color=BLUE, size='12'):
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    element = OxmlElement('w:bottom')
    element.set(qn('w:val'), 'single')
    element.set(qn('w:sz'), size)
    element.set(qn('w:space'), '8')
    element.set(qn('w:color'), color)
    borders.append(element)
    properties.append(borders)


def set_cell_text(cell, text, bold=False, color=NAVY):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = 'Aptos'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_title(doc, title, subtitle=None):
    paragraph = doc.add_paragraph()
    paragraph.space_before = Pt(14)
    paragraph.space_after = Pt(8)
    run = paragraph.add_run(title)
    run.font.name = 'Aptos Display'
    run.font.size = Pt(25)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    border_bottom(paragraph, CYAN, '18')
    if subtitle:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(14)
        run = paragraph.add_run(subtitle)
        run.italic = True
        run.font.name = 'Aptos'
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = 'Aptos Display'
    run.font.bold = True
    run.font.size = Pt(16 if level == 1 else 12.5)
    run.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else NAVY)
    if level == 1:
        border_bottom(paragraph, 'BFDBFE', '7')


def add_text(doc, text, emphasis=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    run.font.name = 'Aptos'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    if emphasis == 'note':
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(BLUE)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style='List Bullet')
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.name = 'Aptos'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(NAVY)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style='List Number')
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.name = 'Aptos'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(NAVY)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, NAVY)
        set_cell_text(cell, header, True, WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 0:
                shade(cells[index], 'F8FAFC')
            set_cell_text(cells[index], str(value), False)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_callout(doc, title, body, color='DBEAFE'):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, color)
    cell.text = ''
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(title + ' ')
    title_run.bold = True
    title_run.font.name = 'Aptos'
    title_run.font.size = Pt(10)
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    body_run = paragraph.add_run(body)
    body_run.font.name = 'Aptos'
    body_run.font.size = Pt(10)
    body_run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, '111827')
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(code)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.2)
    run.font.color.rgb = RGBColor.from_string('E5E7EB')
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('AI Engineer Studio  •  ')
    run.font.name = 'Aptos'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


def document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    add_page_number(section)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.space_before = Pt(72)
    run = paragraph.add_run('AI ENGINEER STUDIO')
    run.font.name = 'Aptos Display'
    run.font.size = Pt(31)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Архитектура, эксплуатация и руководство\nпо созданию аналогичного AI-агента')
    run.font.name = 'Aptos Display'
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(22)
    run = paragraph.add_run('Техническая документация • редакция 1.0\nПодготовлено по фактическим исходникам проекта')
    run.font.name = 'Aptos'
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_callout(doc, 'Назначение документа.', 'Единый справочник для разработчика, владельца продукта и команды сопровождения: устройство приложения, функции, контракты, безопасность, текущие ограничения и пошаговый план реализации AI IDE с нуля.', 'E0F2FE')
    doc.add_page_break()

    add_title(doc, '1. Краткое описание', 'AI Engineer Studio — Windows desktop IDE для диалога с OpenAI-compatible моделями и работы с контекстом локального проекта.')
    add_text(doc, 'Приложение построено как WPF-клиент на .NET 10. Пользователь выбирает или создаёт проект, ведёт диалог с моделью, передаёт в запрос фрагменты файлов и получает Markdown-ответ. Архитектура включает локальный Python bridge для обращения к OpenAI-compatible API и набор инструментов агента для работы с рабочей папкой.')
    add_table(doc, ['Параметр', 'Значение'], [
        ('Платформа', 'Windows desktop'), ('UI', 'WPF, borderless window, собственный titlebar'), ('Клиентский язык', 'C# / .NET 10'), ('AI-транспорт', 'Python 3 + пакет openai'), ('API', 'OpenAI-compatible Chat Completions'), ('Хранение', 'JSON: настройки пользователя и документы проектов'), ('Защита токена', 'Windows DPAPI, область CurrentUser')])

    add_heading(doc, '2. Возможности приложения')
    add_bullets(doc, [
        'Создание и открытие рабочих проектов с именем, папкой, описанием и сохранённой историей диалога.',
        'Просмотр дерева проекта и выбор файлов для передачи в контекст модели.',
        'Многострочный чат с отправкой кнопкой или сочетанием Ctrl+Enter.',
        'Отображение ответов в локальном Markdown renderer: заголовки, жирный текст, списки и fenced-блоки кода.',
        'Настройка API-токена, Base URL, модели, temperature, лимита токенов, системного промпта и темы.',
        'Темы Graphite, Светлая, Midnight Blue и Emerald Night; быстрая смена Graphite/Светлая.',
        'Inspector с зонами Agent Actions, Diff Viewer и Live Preview.',
        'Сервисные классы для работы с файлами, сохранением данных, Git-версией и ограниченным desktop-управлением.',
        'Радужная анимация ожидания ответа и блокировка повторной отправки на уровне UI/состояния запроса.'
    ])

    add_heading(doc, '3. Архитектура и компоненты')
    add_text(doc, 'Решение разделено на слой представления, модели, сервисы и Python-мост. WPF отвечает за состояние интерфейса и UX, сервисы — за файловые операции и хранение, bridge — за AI-вызов и цикл инструментов.')
    add_code(doc, 'Пользователь → MainWindow (WPF) → NeuralBridgeService → tools/neural_bridge.py\n                                            ↓                     ↓\n                                    ProjectFileService      OpenAI-compatible API\n                                            ↓                     ↓\n                                  ProjectStore / SettingsStore ← tool results')
    add_table(doc, ['Слой / файл', 'Ответственность'], [
        ('App.xaml', 'Глобальная палитра DynamicResource и шаблоны WPF-контролов.'),
        ('MainWindow.xaml', 'Структура окна: titlebar, activity bar, explorer, чат, inspector, status bar.'),
        ('MainWindow.xaml.cs', 'Состояние UI, темы, история, контекст, вызов bridge, preview/diff.'),
        ('Models/ChatModels.cs', 'DTO ChatMessage и NeuralRequest.'),
        ('Models/ProjectDocument.cs', 'Модель сохраняемого документа проекта.'),
        ('Services/ProjectFileService.cs', 'Рекурсивный поиск файлов и формирование текстового контекста.'),
        ('Services/SettingsStore.cs', 'JSON-настройки и DPAPI-шифрование API-токена.'),
        ('Services/ProjectStore.cs', 'CRUD JSON-документов проектов в каталоге Projects.'),
        ('Services/NeuralBridgeService.cs', 'Запуск Python-процесса и передача JSON через stdin/stdout.'),
        ('UI/MarkdownRenderer.cs', 'Локальный рендер базового Markdown в WPF-элементы.'),
        ('neural_bridge.py', 'API-клиент, schemas инструментов, проверка путей, agent/tool loop.')])

    add_heading(doc, '4. Данные, хранение и безопасность')
    add_table(doc, ['Объект', 'Расположение / формат', 'Особенности'], [
        ('Настройки', '%AppData%\\AIEngineerStudio\\settings.json', 'API-ключ сохраняется с DPAPI CurrentUser.'),
        ('Проекты', 'Projects\\<имя>.json рядом с приложением', 'Имя, рабочая папка, описание, UpdatedAt, история чата.'),
        ('Контекст', 'Формируется в памяти перед запросом', 'До 15 файлов, до 12 000 символов на файл.'),
        ('Bridge-запрос', 'JSON в stdin дочернего Python-процесса', 'Содержит ключ, параметры модели, историю, системный промпт.'),
        ('Ответ модели', 'stdout bridge', 'Должен соответствовать контракту, который читает клиент.')])
    add_callout(doc, 'Безопасность.', 'DPAPI защищает токен на диске, но ключ передаётся bridge-процессу во время запроса. Это допустимо для локального прототипа, но для production предпочтителен нативный .NET-клиент или отдельный защищённый локальный сервис.', 'FEF3C7')

    add_heading(doc, '5. Алгоритм обработки запроса')
    add_numbered(doc, [
        'Пользователь открывает проект и выбирает файлы в Solution Explorer.',
        'ProjectFileService.ReadContext читает не более 15 файлов, исключая .git, bin, obj, .vs, node_modules и __pycache__.',
        'MainWindow формирует NeuralRequest: ключ, endpoint, модель, параметры генерации, промпт, контекст и историю.',
        'NeuralBridgeService запускает python neural_bridge.py как дочерний процесс и записывает JSON в stdin.',
        'Bridge создаёт OpenAI-клиент, добавляет системные инструкции и передаёт истории модели.',
        'Если модель вызывает инструмент, bridge валидирует путь внутри рабочей папки, выполняет операцию, добавляет tool-result в историю и повторяет запрос.',
        'Когда модель возвращает финальный текст, результат должен быть передан WPF-клиенту и отрисован MarkdownRenderer.',
        'UI разблокирует ввод, сохраняет историю в ProjectStore и выключает индикатор ожидания.'
    ])

    add_heading(doc, '6. Инструменты нейросети')
    add_text(doc, 'Каталог функций описан в SCHEMAS neural_bridge.py. Все файловые пути нормализуются функцией safe(base, value): выход за границы рабочей папки запрещён.')
    add_table(doc, ['Инструмент', 'Назначение', 'Риск / рекомендуемая защита'], [
        ('list_files', 'Вывод элементов каталога.', 'Только чтение.'),
        ('get_project_tree', 'Рекурсивное дерево проекта.', 'Исключать служебные и крупные каталоги.'),
        ('read_file', 'Чтение диапазона строк UTF-8 файла.', 'Лимитировать размер и типы файлов.'),
        ('search_code', 'Regex-поиск по файлам.', 'Лимитировать результаты и время.'),
        ('create_folder', 'Создание папки.', 'Требовать approval для изменений.'),
        ('write_file', 'Создание/перезапись файла.', 'Diff + явное подтверждение пользователя.'),
        ('edit_file_part', 'Точечная замена фрагмента.', 'Diff + подтверждение; проверка уникальности.'),
        ('execute_command', 'Команда оболочки в папке проекта.', 'Allowlist, approval, timeout, журналирование.'),
        ('screenshot', 'Снимок рабочего стола.', 'Согласие пользователя и безопасное хранение.'),
        ('mouse_position', 'Чтение координат мыши.', 'Минимальный риск.'),
        ('move_mouse', 'Перемещение курсора.', 'Только при computer_control_enabled.'),
        ('mouse_click', 'Левый/правый клик.', 'Только при computer_control_enabled и с подтверждением.')])
    add_callout(doc, 'Ключевой принцип.', 'В текущем bridge операции записи и команд выполняются сразу после tool call. Для безопасного агента production-уровня нужен approval flow: план → diff/команда → подтверждение → выполнение → журнал.', 'FEE2E2')

    add_heading(doc, '7. Контракты и важное замечание по текущей реализации')
    add_text(doc, 'В исходниках обнаружено несоответствие протокола: NeuralBridgeService ожидает единый JSON-объект с полями ok и content. neural_bridge.py в текущей редакции печатает line-delimited события вида tool_start/tool_end/final, а финальное событие имеет поле content без ok. Такой stdout невозможно распарсить текущим JsonDocument.Parse как ожидаемый объект.')
    add_code(doc, 'Ожидание C#: { "ok": true, "content": "..." }\nФактический финал bridge: { "event": "final", "content": "..." }')
    add_bullets(doc, [
        'Вариант A: изменить Python bridge, чтобы в конце печатать единый совместимый JSON; tool events писать в stderr или отключить.',
        'Вариант B: изменить C# на построчное чтение stdout, обрабатывать tool_start/tool_end для Inspector и извлекать content из event=final.',
        'Рекомендуется вариант B: он сохраняет наблюдаемость работы агента и позволяет показывать события в UI.'
    ])

    add_heading(doc, '8. Известные ограничения и план улучшений')
    add_table(doc, ['Приоритет', 'Ограничение', 'Рекомендация'], [
        ('Критичный', 'Несовместимый stdout-контракт C# и bridge.', 'Унифицировать протокол событий и покрыть integration-тестом.'),
        ('Критичный', 'Нет полноценного approval flow для write/execute.', 'Добавить очередь согласований, diff и аудит.'),
        ('Высокий', 'Нет надёжной отмены на всём пути запроса.', 'Передавать cancellation в bridge/API и выводить Stop.'),
        ('Высокий', 'История диалога растёт без сжатия.', 'Окно последних сообщений + summary старого контекста.'),
        ('Средний', 'Markdown ограничен.', 'Добавить ссылки, таблицы, sanitization и подсветку кода.'),
        ('Средний', 'Preview демонстрационный.', 'WebView2 sandbox, lifecycle dev server и HMR.'),
        ('Низкий', 'Нет тестов и installer.', 'Unit/integration tests, self-contained publish или MSIX.')])

    doc.add_page_break()
    add_title(doc, '9. Руководство: аналогичный AI-агент с нуля', 'Практическая последовательность для создания безопасного desktop coding agent от прототипа до результата.')
    add_heading(doc, '9.1. Этап 1 — требования и границы')
    add_numbered(doc, [
        'Определите сценарии: чат по коду, анализ, правки, запуск тестов, preview, управление браузером.',
        'Определите trust boundary: рабочая папка, разрешённые расширения, сети, команды и чувствительные данные.',
        'Разделите действия на read-only, reversible write и high-risk execute/desktop control.',
        'Зафиксируйте UX согласования: что выполняется автоматически, а что требует явного клика пользователя.',
        'Сразу задайте бюджеты: максимальные токены, размер контекста, число шагов, таймауты и лимиты стоимости.'
    ])
    add_heading(doc, '9.2. Этап 2 — каркас решения')
    add_bullets(doc, [
        'Создайте WPF/.NET приложение и выделите Models, Services, UI и Views.',
        'Опишите DTO: ChatMessage, AgentRequest, AgentEvent, ToolCall, ToolResult, ApprovalRequest.',
        'Добавьте SettingsStore и защищённое хранение секретов: DPAPI для desktop или OS keychain.',
        'Сделайте ProjectStore с явной схемой JSON и валидацией имён файлов.',
        'Реализуйте ProjectFileService с игнор-листом, лимитами размера и кодировок.'
    ])
    add_heading(doc, '9.3. Этап 3 — AI-транспорт и event protocol')
    add_text(doc, 'Используйте единый JSON Lines-протокол: один объект на строку. Он удобен для streaming, inspector и диагностики. Каждое событие имеет тип, идентификатор сессии, время и полезную нагрузку.')
    add_code(doc, '{"type":"status","message":"Подготавливаю контекст"}\n{"type":"tool_proposed","id":"t-17","name":"edit_file_part","args":{...}}\n{"type":"approval_required","toolCallId":"t-17","preview":"diff ..."}\n{"type":"tool_result","toolCallId":"t-17","success":true,"result":"..."}\n{"type":"final","content":"## Готово\n..."}')
    add_bullets(doc, [
        'Клиент читает stdout построчно асинхронно, десериализует AgentEvent и обновляет UI через Dispatcher.',
        'stderr используется для технических логов, но не смешивается с протоколом.',
        'Любая ошибка также является структурированным событием error с кодом и безопасным сообщением.',
        'Сессия отменяется CancellationToken: останавливается процесс, HTTP-запрос и ожидающие approval.'
    ])
    add_heading(doc, '9.4. Этап 4 — tool loop и безопасность')
    add_numbered(doc, [
        'Модель возвращает tool call; сервер/bridge проверяет имя функции и JSON-аргументы по схеме.',
        'Path guard получает абсолютный путь, canonicalize/resolve и проверяет принадлежность рабочему корню.',
        'Для write/edit генерируется preview diff; для shell-команды — отображается точная команда, cwd и timeout.',
        'Клиент создаёт ApprovalRequest. Без подтверждения tool не исполняется.',
        'После одобрения инструмент запускается в изолированном контексте с timeout и лимитами вывода.',
        'Результат и хеш изменённых файлов записываются в audit log и передаются модели.',
        'Цикл завершается по final-ответу либо контролируемому лимиту шагов; бесконечный loop запрещён.'
    ])
    add_heading(doc, '9.5. Этап 5 — интерфейс и UX')
    add_bullets(doc, [
        'Показывайте состояние: сбор контекста, ожидание сети, tool proposal, ожидание подтверждения, выполнение, финал.',
        'Во время запроса блокируйте повторную отправку и предусмотрите заметную кнопку Stop.',
        'Для правок показывайте side-by-side или unified diff, кнопки «Применить», «Отклонить», «Применить всё».',
        'Для команд показывайте stdout/stderr, код завершения и возможность повторного запуска.',
        'Не используйте только цвет для критических статусов: добавляйте текст, иконку и доступные tooltip.'
    ])
    add_heading(doc, '9.6. Этап 6 — проверка и выпуск')
    add_table(doc, ['Вид проверки', 'Минимальный набор'], [
        ('Unit', 'Path guard, сериализация событий, context limits, DPAPI adapter, diff generator.'),
        ('Integration', 'Mock OpenAI: final, tool call, ошибка, отмена, несколько tool calls.'),
        ('Security', 'Path traversal, shell injection, secrets in logs, неподтверждённая запись.'),
        ('UX', 'Повторная отправка, Stop, недоступная сеть, пустой ответ, длинный Markdown.'),
        ('Release', 'Проверка Python/runtime, diagnostic screen, installer/MSIX, crash logging.')])
    add_callout(doc, 'Definition of Done.', 'Агент считается готовым не после первого успешного ответа, а когда воспроизводимо обрабатывает отмену, ошибки API, подтверждения, лимиты, журналирование и безопасные границы рабочей папки.', 'DCFCE7')

    add_heading(doc, '10. Быстрый запуск текущего проекта')
    add_code(doc, 'cd "AI Engineer Studio WPF"\npython -m pip install openai\npython -m py_compile neural_bridge.py\ndotnet build AIEngineerStudio.Wpf.csproj\ndotnet run --project AIEngineerStudio.Wpf.csproj')
    add_text(doc, 'В настройках приложения укажите API-токен, Base URL, модель и системный промпт. Токен будет сохранён в профиле текущего пользователя. Перед использованием операций изменения файлов или управления компьютером убедитесь, что approval-механизм реализован и включён осознанно.')

    add_heading(doc, '11. Чек-лист для команды')
    add_bullets(doc, [
        'Согласовать единый JSONL-контракт между WPF и Python bridge.',
        'Реализовать approval flow до включения опасных инструментов в production.',
        'Добавить UI для PythonPath и диагностическую проверку зависимостей.',
        'Реализовать отмену генерации и корректное завершение процесса bridge.',
        'Ограничить/сжимать историю и формировать релевантный контекст.',
        'Добавить автоматические тесты и журналирование без утечки секретов.',
        'Подготовить пакетирование приложения и bootstrap Python-зависимостей.'
    ])

    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    document()
